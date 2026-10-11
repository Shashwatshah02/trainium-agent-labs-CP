#!/usr/bin/env python3
"""
generate_word_report.py
Generates a complete, professionally styled Word document (.docx) covering
the entire Hack the Chip hackathon project from start to finish.
"""

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="F0F4F8", border_hex="1F497D"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="36" w:space="0" w:color="{border_hex}"/><w:top w:val="none"/><w:right w:val="none"/><w:bottom w:val="none"/></w:tcBorders>')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_title = p.add_run(f"📌 {title}\n")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(11)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    
    run_text = p.add_run(text)
    run_text.font.name = "Arial"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(0x2A, 0x2A, 0x2A)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def style_heading(p, font_size=16, color_rgb=(0x1F, 0x49, 0x7D), space_before=16, space_after=6):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    for r in p.runs:
        r.font.name = "Arial"
        r.font.size = Pt(font_size)
        r.font.bold = True
        r.font.color.rgb = RGBColor(*color_rgb)

def add_body_p(doc, text="", bold_prefix="", space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = "Arial"
        r_prefix.font.size = Pt(10.5)
        r_prefix.font.bold = True
        r_prefix.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
    if text:
        r_text = p.add_run(text)
        r_text.font.name = "Arial"
        r_text.font.size = Pt(10)
        r_text.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def main():
    doc = Document()
    
    # Page setup
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)

    # ========================== TITLE BLOCK ==========================
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(12)
    title_p.paragraph_format.space_after = Pt(4)
    r_title = title_p.add_run("Autonomous Multi-Engine Overlap Agent for AWS Trainium")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(0x11, 0x2A, 0x4A)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(14)
    r_sub = sub_p.add_run("Complete Project Engineering Master Report & Technical Architecture Guide\n")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(13)
    r_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    
    meta_p = doc.add_paragraph()
    meta_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_p.paragraph_format.space_after = Pt(20)
    r_meta = meta_p.add_run("Event: NYU × Annapurna Labs Hack the Chip 2026 | Team CP: Shashwat Shah, Heet Mehta, Tanay Doijode\nTarget Hardware: AWS Trainium (NeuronCore-v2, Pod seat-140) | LLM: Qwen/Qwen3-8B")
    r_meta.font.name = "Arial"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(0x77, 0x77, 0x77)

    # Divider line
    p_div = doc.add_paragraph()
    p_div.paragraph_format.space_after = Pt(14)
    r_div = p_div.add_run("―" * 58)
    r_div.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)

    # ========================== EXECUTIVE SUMMARY ==========================
    h1 = doc.add_paragraph("1. Executive Summary: What We Built")
    style_heading(h1, font_size=16)

    add_body_p(doc, 
        "Modern deep learning workloads running on AWS Trainium (NeuronCore-v2) suffer from severe hardware underutilization. "
        "Although Trainium silicon features industry-leading raw compute capacity (powering state-of-the-art models like Anthropic Claude), "
        "real-world AI kernels spend 66% to 90% of their execution time sitting completely idle waiting for memory transfers between external "
        "High Bandwidth Memory (HBM) and on-chip SRAM cache (Shared Buffer / SBUF).")

    add_body_p(doc, 
        "To solve this industry-wide memory wall bottleneck, Team CP built the Multi-Engine Overlap Agent: an autonomous AI compiler loop that "
        "analyzes naive sequential code and automatically refactors it into a 3-way parallel hardware pipeline. Our system overlaps Memory Movement "
        "(DMA Engine), Bilinear Interpolation / Scaling (Vector Engine), and Matrix Multiplication (Tensor Engine) so that all three physical engines "
        "run concurrently with zero memory stalls in steady-state.")

    add_callout(doc, "Key Hackathon Results at a Glance",
        "• Hardware Speedup: 2.07× Latency Reduction on standard 8-block workloads (240 µs down to 116 µs).\n"
        "• Peak Speedup: 2.26× Latency Reduction on 16-block stress workloads (480 µs down to 212 µs).\n"
        "• Idle Memory Stalls: Reduced from 66.7% down to 0.0% in steady-state execution.\n"
        "• Mathematical Proof: Roofline Arithmetic Intensity = 30.4 FLOPs/Byte (7.3× below Trainium's 222.0 ridge point, proving 100% memory-boundedness).\n"
        "• Autonomous Convergence: Qwen3-8B converged to a perfect 1.00/1.00 score in just 1 round on live Trainium silicon (pod seat-140).\n"
        "• Numerical Accuracy: 0.00e+00 discrepancy against 64-bit NumPy ground truth across all 7 hostile test cases.")

    # ========================== THE PROBLEM & ORIGIN ==========================
    h2 = doc.add_paragraph("2. The Origin Story & The Silicon Reality")
    style_heading(h2, font_size=16)

    add_body_p(doc, 
        "During our project discovery meeting, the Annapurna Labs Principal Software Engineer provided critical architectural guidance: "
        "Trainium chips have extraordinary mathematical compute capabilities, but the overwhelming bottleneck in production is I/O-bound data movement. "
        "Specifically, kernels frequently perform redundant memory fetches or force the compute engines to stall while waiting for DMA transfers.",
        "The Principal Engineer's Guidance: ")

    add_body_p(doc, 
        "Every NeuronCore-v2 processor contains two physically separated memory tiers:",
        "The Memory Hierarchy: ")
    
    add_body_p(doc, 
        "1. External High Bandwidth Memory (HBM): Large capacity (up to 32GB per device), but physically distant from the core with high latency.\n"
        "2. On-Chip SRAM Shared Buffer (SBUF): Extremely high bandwidth and zero-wait-state access, but tightly capacity-constrained (~24MB total on chip).\n"
        "Crucial Hardware Rule: The math engines (Vector and Tensor) physically CANNOT read or write data directly to HBM. All inputs must be transferred "
        "into SBUF first via the DMA engine before any arithmetic can execute.")

    add_body_p(doc, 
        "Every NeuronCore-v2 core houses three physically distinct, asynchronous execution engines:",
        "The 3 Independent Physical Engines: ")

    add_body_p(doc, 
        "• DMA Engine: Handles asynchronous bidirectional memory transfers between HBM and SBUF (via nisa.dma_copy or nl.load).\n"
        "• Vector Engine: Specialized for 1D and 2D element-wise arithmetic, activations, and bilinear scaling (via nisa.tensor_scalar, nl.add, nl.multiply).\n"
        "• Tensor Engine: Specialized systolic array hardware for high-throughput matrix multiplication (via nisa.nc_matmul).")

    add_body_p(doc, 
        "When an engineer or a general-purpose LLM writes kernel code, it executes sequentially: Block 0 is transferred by DMA (while Vector and Tensor sit idle); "
        "then Block 0 is scaled by Vector (while DMA and Tensor sit idle); then Block 0 is multiplied by Tensor (while DMA and Vector sit idle). "
        "Because two out of the three engines are doing zero work on any given cycle, the chip suffers a permanent 66.7% hardware idle stall (2/3 wasted time).",
        "The 66.7% Naive Idle Stall Disaster: ")

    # ========================== THE SOLUTION: 3-WAY OVERLAP ==========================
    h3 = doc.add_paragraph("3. The Solution: 3-Way Overlapped Engine Pipelining")
    style_heading(h3, font_size=16)

    add_body_p(doc, 
        "Rather than executing sequentially, our system structures kernel execution as a 3-stage software pipeline running concurrently:",
        "The 3-Engine Assembly Line: ")

    add_body_p(doc, 
        "• Stage 1 (Prologue): DMA prefetches Block 0 and Block 1 into SBUF; Vector engine begins scaling Block 0.\n"
        "• Stage 2 (Steady-State Overlap): On every iteration b, all three engines fire in parallel:\n"
        "    - DMA Engine loads Block b+1 into SBUF cache\n"
        "    - Vector Engine scales Block b in SBUF cache\n"
        "    - Tensor Engine computes matrix multiplication on Block b-1 in SBUF cache\n"
        "• Stage 3 (Epilogue Drain): Once all HBM data has been prefetched, the pipeline drains remaining Vector and Tensor stages for the final two blocks.")

    add_body_p(doc, 
        "Because on-chip memory is strictly budgeted on Trainium, buffer management requires precise mathematical constraints:",
        "SBUF Cache Geometry & Triple-Buffering: ")

    add_body_p(doc, 
        "• Partition Dimension Limit (Pmax): Hardware maximum is 128 rows. Tiles cannot exceed 128 rows without causing silicon execution faults.\n"
        "• Free Dimension Limit (Fmax): Hardware maximum is 512 columns.\n"
        "• Tile Selection: We selected 128 × 128 tiles, perfectly matching Trainium systolic array register geometry.\n"
        "• Why Triple-Buffering: Double-buffering (ping-pong) only allows 2 operations to overlap. Because Trainium has 3 concurrent engines, we allocate "
        "three distinct buffers in SBUF: buf_dma, buf_vec, and buf_tensor.\n"
        "• Zero-Copy Pointer Rotation: Instead of copying memory between buffers (which burns memory bandwidth), pointers are rotated in a single CPU cycle:\n"
        "    buf_tensor, buf_vec, buf_dma = buf_vec, buf_dma, buf_tensor (0 µs overhead, 0 bytes moved).")

    # ========================== TEAM BIFURCATION ==========================
    h4 = doc.add_paragraph("4. Team Work Distribution & Strict API Contracts")
    style_heading(h4, font_size=16)

    add_body_p(doc, 
        "To enable simultaneous, conflict-free development across three engineers, the architecture was split into three modular tracks with strict interface contracts:",
        "The 3-Track Modular Architecture: ")

    # Table of tracks
    t_tracks = doc.add_table(rows=4, cols=4)
    t_tracks.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_tracks.autofit = False

    headers = ["Track Lead", "Git Branch", "Main Files", "Core Responsibilities"]
    col_widths = [Inches(1.2), Inches(1.3), Inches(1.8), Inches(2.2)]

    for i, h in enumerate(headers):
        cell = t_tracks.cell(0, i)
        cell.width = col_widths[i]
        set_cell_background(cell, "1F497D")
        set_cell_margins(cell, 120, 120, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_tracks = [
        ("Shashwat Shah\n(Track 1)", "shashwat/grader", "overlap_bench.py\nrun_local_demo.py", "Hardware Grader, AST Code Bouncer, 7 Hostile Test Permutations, Step A-B-C Block-n Isolation Diagnostic."),
        ("Heet Mehta\n(Track 2)", "heet/agent", "agent.py\ndashboard.html\nserve_dashboard.py", "Autonomous Agent Loop, SBUF Geometry Prompt Engineering, Qwen3-8B live silicon deployment, SSE Web Console."),
        ("Tanay Doijode\n(Track 3)", "tanay/benchmark", "reference_pipeline.py\nvisualize_pipeline.py\nSLIDES_OUTLINE.md", "Golden Reference Kernels, Roofline Model (30.4 FLOPs/Byte), High-res Gantt charts, Speedup Benchmarking.")
    ]

    for row_idx, data in enumerate(data_tracks, start=1):
        bg = "FFFFFF" if row_idx % 2 != 0 else "F7FAFC"
        for col_idx, text in enumerate(data):
            cell = t_tracks.cell(row_idx, col_idx)
            cell.width = col_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 100, 100)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ========================== TRACK 1: SHASHWAT'S GRADER ==========================
    h5 = doc.add_paragraph("5. Track 1 Deep Dive: Hardware Grader & Block-n Hazard Diagnostic (Shashwat)")
    style_heading(h5, font_size=16)

    add_body_p(doc, 
        "General-purpose LLMs fail repeatedly at low-level hardware kernel programming because compiler errors are opaque. "
        "To enable autonomous self-healing, Shashwat designed and built a two-tier grading and diagnostic system in overlap_bench.py:",
        "The Grader Architecture: ")

    add_body_p(doc, 
        "1. Static AST Security Bouncer (Pre-Execution Verification):\n"
        "   - Parses code via Python's ast module in under 2 milliseconds without executing unsafe code.\n"
        "   - Banned Cheats: Detects and rejects whole-array shortcuts (np.einsum, matmul, dot, tensordot, kron) and raw whole-array @ operations (x @ weight).\n"
        "   - External Framework Ban: Blocks import torch, scipy, sklearn, and jax, forcing the model to generate native NKI/Trainium primitives.\n"
        "   - SBUF Tile Bound Enforcement: Rejects any tile allocation exceeding Trainium's hardware limit Pmax = 128 rows.\n"
        "   - Engine Verification: Confirms presence of DMA transfers, Vector scaling, Tensor matmul, and multi-buffer pointer rotation.")

    add_body_p(doc, 
        "2. The 7 Hostile Stress Test Permutations:\n"
        "   - Standard Even (256 × 128): Clean 2-block baseline.\n"
        "   - Prime Rows / Ragged Tail (317 × 128): 317 cannot be divided by 128 (128 + 128 + 61 rows). Exposes missing boundary clamps.\n"
        "   - Sub-Tile Matrix (64 × 128): Smaller than a full 128-row partition. Tests single partial-block handling.\n"
        "   - Large Stress Matrix (1024 × 128): 8 full blocks. Verifies sustained multi-buffering with zero memory leaks.\n"
        "   - Extreme Magnitude (10^4 floats): Stresses accumulator overflow and forces float32 intermediate precision before bfloat16 conversion.\n"
        "   - Zero Matrix with Bias: Tests additive scalar bias beta handling when input data is zero.\n"
        "   - Negative Signed Matrix: Verifies proper signed arithmetic across Vector scaling and Tensor matrix multiplications.")

    add_body_p(doc, 
        "3. Step A → B → C Block-n Isolation Diagnostic (The 'Doctor'):\n"
        "   When a candidate kernel fails at Block n, generic unit tests only report a value mismatch. Our diagnostic isolates the failure:\n"
        "   - Step A (Granular Block Isolation): Slices out the exact input block where divergence occurred.\n"
        "   - Step B (Single-Block Micro-Test): Runs that single block through the candidate kernel in complete isolation.\n"
        "   - Step C (Hazard Categorization & Surgical Instruction):\n"
        "     • Passes alone, but fails in pipeline? -> PIPELINE_SYNC_HAZARD (Buffer Overwrite). DMA overwrote SBUF cache before Tensor finished reading.\n"
        "       Surgical Hint: 'Swap active and next buffer pointers explicitly (active_buf, next_buf = next_buf, active_buf)'.\n"
        "     • Fails on final odd-sized tile? -> RAGGED_EDGE_HAZARD. Read past buffer boundary.\n"
        "       Surgical Hint: 'Apply boundary clamp valid = min(128, H - r_start)'.\n"
        "     • Diverges numerically / NaN? -> NUMERICAL_PRECISION_HAZARD. Intermediate overflow.\n"
        "       Surgical Hint: 'Accumulate intermediate products in float32 SBUF before bfloat16 cast'.\n"
        "     • All zeros on a block? -> UNCOMPUTED_BLOCK_HAZARD. Loop condition terminated early.\n"
        "       Surgical Hint: 'Handle single blocks, prefetch in prologue, and drain all remaining blocks in epilogue'.")

    # ========================== TRACK 2: HEET'S AGENT ==========================
    h6 = doc.add_paragraph("6. Track 2 Deep Dive: Autonomous Recalibration Agent (Heet)")
    style_heading(h6, font_size=16)

    add_body_p(doc, 
        "Heet built the autonomous agent driver in agent.py, connecting directly to the Qwen/Qwen3-8B model served on Trainium silicon (pod seat-140):",
        "Agent Loop Architecture: ")

    add_body_p(doc, 
        "• Token Budget Management: Compact SBUF geometry prompt with max_tokens=1000 capped generation time to ~33 seconds per round (a 3.3× speedup).\n"
        "• Progressive Surgical Injection: When the Grader diagnoses a hazard, the agent DOES NOT tell the model to rewrite the code from scratch. "
        "It injects the 1-sentence surgical fix into the prompt, keeping working code structures intact.\n"
        "• Live Silicon Convergence Proof (The Hackathon Milestone):\n"
        "  - Round 0: Candidate kernel skipped computing Block 1 due to a loop boundary condition bug. Grader diagnosed UNCOMPUTED_BLOCK_HAZARD (Score: 0.30).\n"
        "  - Round 1: Model ingested the surgical directive, restructured prologue and epilogue, and converged to a perfect Score: 1.00 / 1.00!\n"
        "  - Verification: 100% pass across all 7 hostile test cases with 100% engine concurrency.\n"
        "• Artifact Ledger (overlap_attempts.jsonl): Automatically archives every attempt, timestamp, token count, latency, score, and code snapshot.")

    # ========================== TRACK 3: TANAY'S BENCHMARKS ==========================
    h7 = doc.add_paragraph("7. Track 3 Deep Dive: Roofline Math & Hardware Acceleration (Tanay)")
    style_heading(h7, font_size=16)

    add_body_p(doc, 
        "Tanay established the mathematical ground truth in reference_pipeline.py and built the hardware profiler and visualizer in visualize_pipeline.py:",
        "Profiling & Empirical Proof: ")

    add_body_p(doc, 
        "1. The Mathematical Roofline Model (1024 × 128 Workload):\n"
        "   - Total Computation: 33,816,576 FLOPs (Bilinear scaling + Tile Matrix Multiplication).\n"
        "   - Total HBM Memory Traffic: 1,114,112 Bytes (~1.08 MB).\n"
        "   - Arithmetic Intensity: 33,816,576 / 1,114,112 = 30.4 FLOPs/Byte.\n"
        "   - AWS Trainium Ridge Point: 222.0 FLOPs/Byte.\n"
        "   - Hardware Classification: 100% MEMORY_BOUND (7.3× below ridge point). Proves mathematically that memory latency dominates runtime.")

    add_body_p(doc, 
        "2. Hardware Speedup Benchmarks Across Workload Dimensions:",
        "Empirical Latency Table: ")

    # Table of speedup
    t_speedup = doc.add_table(rows=4, cols=6)
    t_speedup.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_speedup.autofit = False

    s_headers = ["Matrix Shape", "Blocks", "Sequential Latency", "Overlapped Latency", "Hardware Speedup", "Memory Stall Reduction"]
    s_widths = [Inches(1.1), Inches(0.7), Inches(1.2), Inches(1.2), Inches(1.1), Inches(1.2)]

    for i, h in enumerate(s_headers):
        cell = t_speedup.cell(0, i)
        cell.width = s_widths[i]
        set_cell_background(cell, "1F497D")
        set_cell_margins(cell, 120, 120, 120, 120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.size = Pt(9)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_speedup = [
        ("512 × 128", "4 Blocks", "120.0 µs", "68.0 µs", "1.76× Speedup", "43.3% Reduction"),
        ("1024 × 128", "8 Blocks", "240.0 µs", "116.0 µs", "2.07× Speedup", "51.7% Reduction"),
        ("2048 × 128", "16 Blocks", "480.0 µs", "212.0 µs", "2.26× Speedup", "55.8% Reduction")
    ]

    for row_idx, data in enumerate(data_speedup, start=1):
        bg = "FFFFFF" if row_idx % 2 != 0 else "F7FAFC"
        for col_idx, text in enumerate(data):
            cell = t_speedup.cell(row_idx, col_idx)
            cell.width = s_widths[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 100, 100, 100, 100)
            p = cell.paragraphs[0]
            if col_idx in (1, 2, 3, 4, 5):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.name = "Arial"
            r.font.size = Pt(9)
            if "Speedup" in text:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)
            else:
                r.font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Embed Gantt chart if present
    gantt_path = os.path.join(os.path.dirname(__file__), "pipeline_gantt_1024.png")
    if os.path.exists(gantt_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        doc.add_picture(gantt_path, width=Inches(6.2))
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run("Figure 1: Hardware Gantt Chart Comparison — Naive Sequential (66.7% idle stalls) vs. 3-Way Overlapped Pipeline (0% memory stalls in steady-state, 2.07× speedup).")
        r_cap.font.name = "Arial"
        r_cap.font.size = Pt(8.5)
        r_cap.font.italic = True
        r_cap.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # ========================== THE WEB DASHBOARD ==========================
    h8 = doc.add_paragraph("8. The Interactive Web Dashboard & Real-Time Console")
    style_heading(h8, font_size=16)

    add_body_p(doc, 
        "To provide a world-class demonstration for hackathon judges, Heet developed a full-stack real-time console:",
        "Web Dashboard Architecture: ")

    add_body_p(doc, 
        "• Frontend (dashboard.html / index.html): A responsive dark-mode cockpit featuring real-time hero metric cards, interactive tabbed navigation "
        "(Overview, Live Execution, Attempt Trajectory, Code Comparison), and embedded SVG/PNG hardware timeline charts.\n"
        "• Backend Server (serve_dashboard.py): A custom lightweight HTTP server running on port 8080.\n"
        "• Live Process Streaming via Server-Sent Events (SSE): When a user clicks 'Run Full Selftest' or 'Run Overlap Agent', the backend spawns a Python "
        "subprocess and streams stdout/stderr line-by-line into the browser terminal in real time.\n"
        "• Dynamic Metrics API (/api/metrics): Synchronizes metrics (FLOPs, Arithmetic Intensity, Latency, and Stall Reductions) dynamically from the codebase.")

    # ========================== LOCAL REPRODUCTION ==========================
    h9 = doc.add_paragraph("9. Reproduction Instructions & Terminal Commands")
    style_heading(h9, font_size=16)

    add_body_p(doc, "All deliverables run seamlessly on both macOS laptops (via NumPy simulation and AST analysis) and physical AWS Trainium silicon pods:")

    add_callout(doc, "Terminal Execution Cheat Sheet",
        "# 1. Run Complete System Demo (Runs all subsystems in ~1 second):\n"
        ".venv/bin/python projects/03-multi-engine-overlap/run_local_demo.py\n\n"
        "# 2. Run Grader & 10/10 Selftest (Shashwat's Track):\n"
        ".venv/bin/python projects/03-multi-engine-overlap/overlap_bench.py --selftest\n\n"
        "# 3. Evaluate Winning Synthesized Kernel:\n"
        ".venv/bin/python projects/03-multi-engine-overlap/eval_best_kernel.py\n\n"
        "# 4. Run Autonomous Agent Recalibration Loop (Offline Replay):\n"
        ".venv/bin/python projects/03-multi-engine-overlap/agent.py --offline --rounds 3\n\n"
        "# 5. Run Profiler & Generate Gantt Chart (Tanay's Track):\n"
        ".venv/bin/python projects/03-multi-engine-overlap/visualize_pipeline.py --rows 1024 --cols 128\n\n"
        "# 6. Launch Live Interactive Web Dashboard:\n"
        ".venv/bin/python projects/03-multi-engine-overlap/serve_dashboard.py\n"
        "# Open browser to: http://localhost:8080")

    # ========================== Q&A DEFENSE GUIDE ==========================
    h10 = doc.add_paragraph("10. Q&A Defense Guide for Annapurna Labs Judges")
    style_heading(h10, font_size=16)

    qa_list = [
        ("Why did you use triple-buffering instead of double-buffering?",
         "Double-buffering only allows two operations to overlap (e.g., DMA and compute). But Trainium has three physically independent engines running simultaneously: DMA, Vector, and Tensor. If you only had two buffers, the third engine would be starved. With triple-buffering (buf_dma, buf_vec, buf_tensor), each engine owns its dedicated on-chip buffer simultaneously, enabling steady-state execution with zero stalls."),
        
        ("How does your Block-n Isolation Diagnostic prove a bug is a memory race condition rather than bad math?",
         "Through our Step B Single-Block Micro-Test. We extract the exact input slice for Block n and run it through the kernel completely alone as an isolated 1-block matrix. If that single block computes the exact correct mathematical answer alone, the algorithmic code is sound. The only reason it produced wrong numbers in the multi-block pipeline is that the DMA engine overwrote the shared SBUF buffer before the Tensor engine finished reading. That proves definitively that it is a Pipeline Sync Hazard."),
        
        ("What happens if input matrix rows aren't a clean multiple of 128?",
         "We built hostile test cases specifically for ragged tails—like our 317 × 128 test, where the last tile has only 61 rows. If the model forgets to clamp the boundary, our diagnostic detects a RAGGED_EDGE_HAZARD and injects valid_rows = min(128, H - r_start) to prevent out-of-bounds reads."),
        
        ("What is your workload's Arithmetic Intensity, and why does pipelining help?",
         "Our workload has an Arithmetic Intensity of 30.4 FLOPs/Byte, which sits 7.3× below Trainium's ridge point of 222.0 FLOPs/Byte. That proves mathematically that the workload is heavily memory-bound. Pipelining allows the DMA Engine to prefetch Block N+1 from HBM into SBUF concurrently while the Vector and Tensor engines compute Block N and N-1, effectively hiding 100% of memory latency behind compute."),
        
        ("Did you sacrifice any mathematical precision to get the 2.07× speedup?",
         "Zero. We achieved exact 0.00e+00 discrepancy against the NumPy 64-bit reference across all 7 test cases. The speedup comes purely from concurrent hardware engine scheduling and SBUF pointer rotation, not from mathematical approximation.")
    ]

    for q, a in qa_list:
        add_body_p(doc, q, "Q: ", space_after=2)
        add_body_p(doc, a, "A: ", space_after=8)

    # Save document
    output_docx = os.path.join(os.path.dirname(__file__), "AWS_Trainium_Project_Master_Report.docx")
    doc.save(output_docx)
    print(f"[SUCCESS] Professionally formatted Word document saved to: {output_docx}")

if __name__ == "__main__":
    main()
